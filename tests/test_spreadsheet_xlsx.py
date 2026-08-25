"""Unit tests for XLSX parser with openpyxl fixtures."""

from __future__ import annotations

import io
import unittest

from openpyxl import Workbook

from documents.models import CELL_NUMBER, CELL_STRING
from documents.parsers.xlsx import XlsxDocumentParser


def _xlsx_bytes(*, cells=None, sheet_title="Sheet1"):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    for coord, value in (cells or {"A1": "hello", "B1": 7}).items():
        ws[coord] = value
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), ws.title


class SpreadsheetXlsxTests(unittest.TestCase):
    def test_parse_xlsx_openpyxl_fixture(self):
        data, sheet_name = _xlsx_bytes(cells={"A1": "Name", "B1": "Age", "A2": "Ada", "B2": 36})
        parsed = XlsxDocumentParser().parse(
            document_id="d-xlsx",
            data=data,
            filename="people.xlsx",
            limits={"max_sheets": 10, "max_table_cells": 10_000},
        )
        self.assertEqual(parsed.parser_id, "xlsx_v1")
        self.assertIsNotNone(parsed.workbook)
        self.assertIn(sheet_name, parsed.workbook.sheet_names)
        self.assertTrue(parsed.cells)
        values = {(c.coordinate, c.value_type, c.value) for c in parsed.cells}
        self.assertIn(("A1", CELL_STRING, "Name"), values)
        self.assertIn(("B2", CELL_NUMBER, "36"), values)
        self.assertTrue(parsed.tables)
        self.assertTrue(parsed.text_blocks)


if __name__ == "__main__":
    unittest.main()
