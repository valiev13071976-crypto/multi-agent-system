"""TCL.xlsx end-to-end defect closure: PRODUCT-FIRST acceptance for the
complete real flow --

    TCL.xlsx (real supplier headers, e.g. "Модель"/"Название", never
    "sku"/"артикул")
        -> understand product identity (Step 2 -- generic role-mapping
           fallback, ``data_intel.mapping``/``data_intel.service``)
        -> check existing products in LIVE Bitrix (Step 3 --
           ``BitrixProductBridge.check_live_existence``, including
           inactive products)
    -> prepare only NEW products (Step 4 -- batch preview,
           ``CHECK_BITRIX_EXISTENCE_BATCH``)
        -> show batch preview (TOTAL/NEW/EXISTING/AMBIGUOUS/INVALID,
           zero writes)
        -> explicit confirmation (``CONFIRM_BATCH_BITRIX_CREATE``)
        -> create selected NEW products inactive (Step 5 -- reuses the
           EXISTING single-product ``execute_single_product_write``
           primitive once per row, never a second write path)
        -> read-back
        -> rerun without duplicates

Three real production defects motivated this closure:

  1. All 34 rows of the real TCL.xlsx were classified INVALID
     (``missing_sku_or_title``) because its headers ("Модель"/
     "Название") had no role mapping at all. Closed by
     ``data_intel.mapping``'s new "модель"/"model"/"код_модели"/
     "код_товара" -> ``ROLE_ARTICLE`` aliases and "название"/
     "наименование" -> ``ROLE_PRODUCT_NAME`` remap (moved off the
     unread ``ROLE_COMPANY_NAME``).
  2. SKU 32LQ63806LC.ARUG was created TWICE in real Bitrix (IDs 994/995)
     because (a) the single-product write's idempotency key was an
     ephemeral per-turn request id instead of the deterministic
     tenant+sku+title default, and (b) neither the single-product nor
     the (then-nonexistent) batch write path ever re-checked the
     CONNECTED (LIVE/FIXTURE) catalog by article/SKU before creating.
     Closed by the idempotency-key fix in
     ``WorkflowPandaConversationGateway._invoke_controlled_bitrix_write``
     plus the new ``BitrixProductBridge.check_live_existence`` guard,
     now consulted before EVERY create (single-product AND batch).
  3. AFTER (1) and (2) were closed, the real TCL.xlsx STILL classified
     all 34 rows INVALID: identity extraction now correctly resolved a
     ``sku`` (e.g. "55C6K") for every row, but the batch existence
     check's own validation additionally required a non-empty
     ``title`` ("if not title or not sku: missing_sku_or_title") --
     and the real TCL.xlsx rows have no usable title column at all.
     Product identity contract fix: a reliable canonical SKU/article/
     model identity is SUFFICIENT on its own for a Bitrix existence/
     duplicate lookup; ``title`` is descriptive metadata, never a
     mandatory uniqueness key. A row is INVALID only when there is
     genuinely no usable identity (no sku AND no title-derived
     fallback) -- never because title alone is absent. Closed in
     ``WorkflowPandaConversationGateway._check_bitrix_existence_batch``
     (validation now checks ``sku`` only) -- the single-product path
     already only keyed its own duplicate guard
     (``BitrixProductBridge.check_live_existence``) by ``sku``, so no
     change was needed there for the two paths to agree.

FIXTURE Bitrix adapter only -- zero live Bitrix mutations anywhere in
this module.
"""

from __future__ import annotations

import io
import os
import unittest
from unittest import mock

from openpyxl import Workbook

from artifacts.service import ArtifactService
from artifacts.store import InMemoryArtifactStore
from business_assistant.action_continuation import (
    CALL_CONTROLLED_BITRIX_WRITE,
    CHECK_BITRIX_EXISTENCE_BATCH,
    CONFIRM_BATCH_BITRIX_CREATE,
    is_batch_bitrix_create_confirmation,
    is_batch_bitrix_existence_check_request,
)
from business_assistant.conversation_gateway import ConversationRequest, WorkflowPandaConversationGateway
from data_intel.contracts import ROLE_ARTICLE, ROLE_PRODUCT_NAME, ROLE_UNKNOWN
from data_intel.mapping import map_header_role
from data_intel.service import DataIntelligenceService
from data_intel.store import InMemoryDatasetStore
from integrations.activation.models import ENV_FIXTURE
from integrations.activation.service import IntegrationActivationService
from integrations.bitrix.catalog import BitrixCatalogStore
from integrations.bitrix.fixture_adapter import BitrixFixtureAdapter
from integrations.bitrix.product_bridge import BitrixProductBridge
from tools.gateway import ToolGateway
from tools.platform.bootstrap import register_platform_tools
from tools.registry import ToolRegistry

TENANT = "tenant-tcl"
OWNER = "user-a"
CONV = "conv-tcl"
FILENAME = "TCL.xlsx"

# Seeded in the shared fixture catalog (integrations/bitrix/catalog.py):
# SKU-X100 (active), SKU-X200 (INACTIVE), SKU-AMBIG (duplicated below to
# force a genuine ambiguous match, same pattern the existing batch-check
# test already uses).
EXISTING_ACTIVE_SKU = "SKU-X100"
EXISTING_INACTIVE_SKU = "SKU-X200"
AMBIGUOUS_SKU = "SKU-AMBIG"
NEW_SKU = "TCL-77Q10K.ARUG"
NEW_TITLE = "TCL 77Q10K QLED телевизор"
SUBSET_NEW_SKU_1 = "55C6K"
SUBSET_NEW_SKU_2 = "65RM7L"

BATCH_CHECK_TEXT = (
    "Проверь весь прайс перед загрузкой на сайт. Покажи, какие товары уже есть в Bitrix, "
    "каких нет и где есть неоднозначность. Ничего не записывай."
)
BATCH_CONFIRM_TEXT = "Подтверждаю: создай эти новые товары в Bitrix."


def _tcl_xlsx_bytes() -> bytes:
    """Real TCL.xlsx-shaped headers: "Модель" (manufacturer model code,
    the identity column) + "Название" (display name) + "Цена" (bare
    generic price) -- deliberately NEVER "sku"/"артикул"/"price", the
    exact real-world shape that used to leave every row INVALID."""
    wb = Workbook()
    ws = wb.active
    ws.append(["Модель", "Название", "Цена"])
    ws.append([NEW_SKU, NEW_TITLE, "54990"])
    ws.append([EXISTING_ACTIVE_SKU, "Samsung Galaxy S24 (из прайса)", "49990"])
    ws.append([EXISTING_INACTIVE_SKU, "Samsung Phone Case (из прайса)", "990"])
    ws.append([AMBIGUOUS_SKU, "Samsung Accessory (из прайса)", "1500"])
    # Genuinely identity-less row: no model code AND no name -- must stay
    # INVALID, never a fabricated identity.
    ws.append(["", "", "1000"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _bitrix_bridge_and_store() -> tuple[BitrixProductBridge, BitrixCatalogStore]:
    store = BitrixCatalogStore()
    # Force a genuine Bitrix-side duplicate for AMBIGUOUS_SKU (same
    # pattern ``tests/test_panda_batch_bitrix_existence_check_defect_
    # closure.py`` and ``tests/test_block5_6_bitrix_aspro_integration.py``
    # already use).
    store.create_product(tenant_id=TENANT, payload={"name": "Dup", "article": AMBIGUOUS_SKU, "price": "10"})
    activation = IntegrationActivationService()
    adapter = BitrixFixtureAdapter(store=store)
    activation._adapters["bitrix"] = adapter  # noqa: SLF001
    ref = activation.put_secret_ref(tenant_id=TENANT, secret_ref="secret:bitrix-tenant-a", value="tok")
    conn = activation.configure_connection(
        tenant_id=TENANT, provider_id="bitrix", credential_ref=ref, environment=ENV_FIXTURE
    )
    activation.verify_connection(tenant_id=TENANT, connection_id=conn.connection_id)
    activation.activate_connection(tenant_id=TENANT, connection_id=conn.connection_id)
    bridge = BitrixProductBridge(integration_activation=activation, environment=ENV_FIXTURE, store=store)
    return bridge, store


def _panda(bridge: BitrixProductBridge) -> tuple[WorkflowPandaConversationGateway, ArtifactService]:
    svc = DataIntelligenceService(InMemoryDatasetStore())
    artifact_service = ArtifactService(store=InMemoryArtifactStore())
    svc.artifact_service = artifact_service
    registry = ToolRegistry()
    register_platform_tools(registry, data_intelligence=svc)
    gateway = ToolGateway(registry=registry, register_search=False)
    panda = WorkflowPandaConversationGateway(
        workflow_engine=object(),
        run_router=object(),
        context_manager=object(),
        tool_gateway=gateway,
        artifact_service=artifact_service,
        bitrix_product_bridge=bridge,
    )
    return panda, artifact_service


async def _register_upload(artifact_service: ArtifactService, *, content: bytes) -> str:
    rec = artifact_service.register_upload(tenant_id=TENANT, owner_id=OWNER, filename=FILENAME, content=content)
    artifact_service.attach_to_conversation(tenant_id=TENANT, artifact_id=rec.artifact_id, conversation_id=CONV)
    return rec.artifact_id


async def _upload_and_analyze(panda: WorkflowPandaConversationGateway, artifact_service: ArtifactService) -> None:
    ref = await _register_upload(artifact_service, content=_tcl_xlsx_bytes())
    await panda.respond(
        ConversationRequest(
            text="Вот прайс-лист TCL, посмотри, что там есть.",
            tenant_id=TENANT,
            user_id=OWNER,
            request_id="r1",
            conversation_id=CONV,
            attachment_refs=(ref,),
        )
    )
    await panda.respond(
        ConversationRequest(
            text="Проанализируй этот прайс.",
            tenant_id=TENANT,
            user_id=OWNER,
            request_id="r2",
            conversation_id=CONV,
        )
    )


class Step2ProductIdentityRoleMappingUnitTests(unittest.TestCase):
    """Mandatory acceptance #1/#2: supplier-style identity headers map to a
    USABLE role, generic enough for any supplier feed (never hardcoded to
    "TCL"/a specific model/row/filename), while a genuinely identity-less
    header still resolves to ROLE_UNKNOWN -- never a fabricated identity."""

    def test_model_code_headers_map_to_article_role(self):
        for header in ("Модель", "модель", "Model", "Код модели", "Код товара"):
            role, _ = map_header_role(header)
            self.assertEqual(role, ROLE_ARTICLE, f"{header!r} must map to ROLE_ARTICLE")

    def test_display_name_headers_map_to_product_name_role(self):
        for header in ("Название", "Наименование", "название", "наименование"):
            role, _ = map_header_role(header)
            self.assertEqual(role, ROLE_PRODUCT_NAME, f"{header!r} must map to ROLE_PRODUCT_NAME")

    def test_genuinely_unrelated_header_stays_unknown(self):
        role, _ = map_header_role("Примечание")
        self.assertEqual(role, ROLE_UNKNOWN)


class BatchExistenceCheckPredicateUnitTests(unittest.TestCase):
    def test_batch_check_request_is_not_a_create_confirmation(self):
        self.assertTrue(is_batch_bitrix_existence_check_request(BATCH_CHECK_TEXT))
        self.assertFalse(is_batch_bitrix_create_confirmation(BATCH_CHECK_TEXT))

    def test_batch_create_confirmation_is_not_an_existence_check(self):
        self.assertTrue(is_batch_bitrix_create_confirmation(BATCH_CONFIRM_TEXT))
        self.assertFalse(is_batch_bitrix_existence_check_request(BATCH_CONFIRM_TEXT))

    def test_singular_write_confirmation_is_not_a_batch_confirmation(self):
        self.assertFalse(is_batch_bitrix_create_confirmation("Подтверждаю: создай этот товар в Bitrix."))


class TclStyleBatchPreviewAcceptanceTests(unittest.IsolatedAsyncioTestCase):
    """Mandatory acceptance #1, #2, #3, #4, #5, #6, #8: the batch preview
    over REAL TCL.xlsx-style headers correctly classifies every row and
    performs zero writes."""

    async def test_batch_preview_classifies_all_five_rows_correctly(self):
        bridge, store = _bitrix_bridge_and_store()
        panda, artifact_service = _panda(bridge)
        await _upload_and_analyze(panda, artifact_service)
        before_catalog_size = len(store.catalog(TENANT))

        result = await panda.respond(
            ConversationRequest(
                text=BATCH_CHECK_TEXT,
                tenant_id=TENANT,
                user_id=OWNER,
                request_id="r3",
                conversation_id=CONV,
            )
        )

        self.assertEqual(result.metadata.get("action_decision"), CHECK_BITRIX_EXISTENCE_BATCH)
        check = result.metadata.get("bitrix_existence_check") or {}
        self.assertEqual(check.get("total"), 5)

        new_skus = {row["sku"] for row in check.get("new") or []}
        existing_skus = {row["sku"] for row in check.get("existing") or []}
        ambiguous_skus = {row["sku"] for row in check.get("ambiguous") or []}
        invalid_reasons = [row["reason"] for row in check.get("invalid") or []]

        # #1 supplier-style headers ("Модель"/"Название") resolved a
        # usable product identity for every row that carries one.
        self.assertEqual(new_skus, {NEW_SKU})
        # #3 an active existing product -> EXISTING.
        # #6 an INACTIVE existing product is STILL detected -> EXISTING.
        self.assertEqual(existing_skus, {EXISTING_ACTIVE_SKU, EXISTING_INACTIVE_SKU})
        # #5 multiple Bitrix matches -> AMBIGUOUS.
        self.assertEqual(ambiguous_skus, {AMBIGUOUS_SKU})
        # #2 the genuinely identity-less row stays INVALID, never a
        # fabricated identity.
        self.assertEqual(len(check.get("invalid") or []), 1)
        self.assertIn("missing_sku_or_title", invalid_reasons)

        # #4 the brand-new product is READY_TO_CREATE.
        ready_rows = [row for row in check.get("new") or [] if row.get("planned_action") == "READY_TO_CREATE"]
        self.assertEqual({row["sku"] for row in ready_rows}, {NEW_SKU})

        # #8 the preview performs ZERO writes.
        self.assertEqual(len(store.catalog(TENANT)), before_catalog_size)
        self.assertEqual(result.metadata.get("mutated"), False)


# Real TCL.xlsx-shaped headers with NO usable title/name column at all --
# the exact production shape ("55C6K", "65RM7L", "75RM7L", ...) that
# still classified every row INVALID after defect (1)/(2) above were
# closed, because the batch check additionally required a non-empty
# ``title``.
TITLE_ABSENT_NEW_SKU = "55C6K"
TITLE_ABSENT_EXISTING_INACTIVE_SKU = EXISTING_INACTIVE_SKU
TITLE_ABSENT_AMBIGUOUS_SKU = AMBIGUOUS_SKU


def _tcl_xlsx_bytes_no_title_column() -> bytes:
    """Real production shape: only a "Модель" identity column and a bare
    "Цена" column -- deliberately NO "Название"/"Наименование" column at
    all, so every row's projected ``title`` is genuinely empty. A strong
    canonical article/model identity (``sku``) must still be sufficient
    on its own for the Bitrix existence/duplicate lookup."""
    wb = Workbook()
    ws = wb.active
    ws.append(["Модель", "Цена"])
    ws.append([TITLE_ABSENT_NEW_SKU, "54990"])
    ws.append([TITLE_ABSENT_EXISTING_INACTIVE_SKU, "990"])
    ws.append([TITLE_ABSENT_AMBIGUOUS_SKU, "1500"])
    # Genuinely identity-less row: no model code at all (and, as in this
    # fixture, no title column either) -- must stay INVALID, never a
    # fabricated identity.
    ws.append(["", "1000"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TitleOptionalProductIdentityAcceptanceTests(unittest.IsolatedAsyncioTestCase):
    """Product identity contract fix (production defect closure after
    #105): a reliable canonical SKU/article/model identity is SUFFICIENT
    on its own for a Bitrix existence/duplicate lookup -- ``title`` is
    descriptive metadata only, never a mandatory uniqueness key. A row is
    INVALID only when there is genuinely no usable identity at all
    (neither sku/article/model). Reproduces the real TCL.xlsx shape (no
    title/name column whatsoever) and proves NEW/EXISTING/AMBIGUOUS
    classification -- never ``missing_sku_or_title`` -- for every row
    that carries a strong model identity, with zero writes."""

    async def _upload_no_title_column(self, panda, artifact_service):
        ref = await _register_upload(artifact_service, content=_tcl_xlsx_bytes_no_title_column())
        await panda.respond(
            ConversationRequest(
                text="Вот прайс-лист TCL, посмотри, что там есть.",
                tenant_id=TENANT,
                user_id=OWNER,
                request_id="r1",
                conversation_id=CONV,
                attachment_refs=(ref,),
            )
        )
        await panda.respond(
            ConversationRequest(
                text="Проанализируй этот прайс.",
                tenant_id=TENANT,
                user_id=OWNER,
                request_id="r2",
                conversation_id=CONV,
            )
        )

    async def test_sku_present_title_absent_is_never_invalid(self):
        bridge, store = _bitrix_bridge_and_store()
        panda, artifact_service = _panda(bridge)
        await self._upload_no_title_column(panda, artifact_service)
        before_catalog_size = len(store.catalog(TENANT))

        result = await panda.respond(
            ConversationRequest(
                text=BATCH_CHECK_TEXT,
                tenant_id=TENANT,
                user_id=OWNER,
                request_id="r3",
                conversation_id=CONV,
            )
        )

        self.assertEqual(result.metadata.get("action_decision"), CHECK_BITRIX_EXISTENCE_BATCH)
        check = result.metadata.get("bitrix_existence_check") or {}
        self.assertEqual(check.get("total"), 4)

        new_skus = {row["sku"] for row in check.get("new") or []}
        existing_skus = {row["sku"] for row in check.get("existing") or []}
        ambiguous_skus = {row["sku"] for row in check.get("ambiguous") or []}
        invalid_rows = check.get("invalid") or []

        # A brand-new sku with NO title at all -> NEW/READY_TO_CREATE,
        # never INVALID: a strong model identity is sufficient on its
        # own, and no title was ever invented to pass validation.
        self.assertEqual(new_skus, {TITLE_ABSENT_NEW_SKU})
        new_row = next(row for row in check.get("new") or [] if row["sku"] == TITLE_ABSENT_NEW_SKU)
        self.assertEqual(new_row.get("title"), "")
        self.assertEqual(new_row.get("planned_action"), "READY_TO_CREATE")

        # An exact INACTIVE existing SKU with no title -> EXISTING (the
        # live duplicate guard is keyed by sku only, matching the batch
        # check).
        self.assertEqual(existing_skus, {TITLE_ABSENT_EXISTING_INACTIVE_SKU})

        # Multiple exact Bitrix matches for the same sku -> AMBIGUOUS.
        self.assertEqual(ambiguous_skus, {TITLE_ABSENT_AMBIGUOUS_SKU})

        # Only the genuinely identity-less row (no sku AND no title) is
        # INVALID -- exactly one row, never all four.
        self.assertEqual(len(invalid_rows), 1)
        self.assertEqual(invalid_rows[0].get("sku"), "")
        self.assertIn("missing_sku_or_title", [row["reason"] for row in invalid_rows])

        # Zero writes during this read-only preview.
        self.assertEqual(len(store.catalog(TENANT)), before_catalog_size)
        self.assertEqual(result.metadata.get("mutated"), False)
        self.assertFalse(result.metadata.get("mutated"))


class BatchLiveConnectionBootstrapAcceptanceTests(unittest.IsolatedAsyncioTestCase):
    """Regression for the production-only failure where every valid TCL row
    became ``bitrix_check_failed`` because the batch read path did not
    bootstrap the per-tenant LIVE Bitrix connection before lookups."""

    async def test_batch_preview_bootstraps_connection_once_before_row_lookups(self):
        bridge, store = _bitrix_bridge_and_store()
        calls: list[str] = []
        original_bootstrap = bridge.ensure_live_connection_ready
        original_lookup = bridge.check_live_existence

        def _bootstrap(*, tenant_id: str) -> None:
            calls.append(f"bootstrap:{tenant_id}")
            original_bootstrap(tenant_id=tenant_id)

        def _lookup(*, tenant_id: str, sku: str, connection_id=None):
            calls.append(f"lookup:{sku}")
            return original_lookup(tenant_id=tenant_id, sku=sku, connection_id=connection_id)

        bridge.ensure_live_connection_ready = _bootstrap
        bridge.check_live_existence = _lookup

        panda, artifact_service = _panda(bridge)
        await _upload_and_analyze(panda, artifact_service)
        before_catalog_size = len(store.catalog(TENANT))

        result = await panda.respond(
            ConversationRequest(
                text=BATCH_CHECK_TEXT,
                tenant_id=TENANT,
                user_id=OWNER,
                request_id="r3-bootstrap",
                conversation_id=CONV,
            )
        )

        check = result.metadata.get("bitrix_existence_check") or {}
        self.assertEqual(calls.count(f"bootstrap:{TENANT}"), 1)
        first_lookup_index = next(i for i, value in enumerate(calls) if value.startswith("lookup:"))
        self.assertLess(calls.index(f"bootstrap:{TENANT}"), first_lookup_index)
        self.assertNotIn("bitrix_check_failed", [row.get("reason") for row in check.get("invalid") or []])
        self.assertEqual(len(store.catalog(TENANT)), before_catalog_size)
        self.assertFalse(result.metadata.get("mutated"))


def _tcl_subset_two_new_no_title_bytes() -> bytes:
    """Two brand-new supplier rows with strong model identity and no title.
    Reproduces the production shape that previously got interpreted as two
    sequential Excel text filters (AND), yielding zero rows."""
    wb = Workbook()
    ws = wb.active
    ws.append(["Модель", "Цена"])
    ws.append([SUBSET_NEW_SKU_1, "54990"])
    ws.append([SUBSET_NEW_SKU_2, "64990"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class DirectMultiSkuSitePreviewAcceptanceTests(unittest.IsolatedAsyncioTestCase):
    def test_workset_uses_current_dataset_id_contract(self):
        """Guard the exact production regression: Workset has
        ``current_dataset_id``; reading a nonexistent ``dataset_id``
        silently disables the entire direct multi-SKU seam."""
        from business_assistant import workset as workset_lib

        ws = workset_lib.start_new_source(
            None,
            tenant_id=TENANT,
            owner_id=OWNER,
            conversation_id=CONV,
            dataset_id="dataset-contract-check",
        )
        self.assertEqual(ws.current_dataset_id, "dataset-contract-check")
        self.assertFalse(hasattr(ws, "dataset_id"))

    """Fresh attachment + two exact SKUs in the same site-preview request
    must stay multi-product and never collapse into Managed Agent's
    legitimate single-product selection semantics.
    """

    async def _assert_direct_multi_preview(self, *, managed_enabled: bool, request_suffix: str):
        bridge, store = _bitrix_bridge_and_store()
        panda, artifact_service = _panda(bridge)
        ref = await _register_upload(artifact_service, content=_tcl_subset_two_new_no_title_bytes())
        before_catalog_size = len(store.catalog(TENANT))

        old_flag = os.environ.get("PANDA_MANAGED_AGENT_ENABLED")
        os.environ["PANDA_MANAGED_AGENT_ENABLED"] = "true" if managed_enabled else "false"
        try:
            result = await panda.respond(
                ConversationRequest(
                    text=(
                        f"Подготовь для сайта только товары {SUBSET_NEW_SKU_1} и {SUBSET_NEW_SKU_2}. "
                        "Покажи полный предпросмотр того, что будет записано в Bitrix. "
                        "Пока ничего не записывай."
                    ),
                    tenant_id=TENANT,
                    user_id=OWNER,
                    request_id=f"direct-multi-{request_suffix}",
                    conversation_id=CONV,
                    attachment_refs=(ref,),
                )
            )
        finally:
            if old_flag is None:
                os.environ.pop("PANDA_MANAGED_AGENT_ENABLED", None)
            else:
                os.environ["PANDA_MANAGED_AGENT_ENABLED"] = old_flag

        self.assertEqual(result.metadata.get("action_decision"), "PREVIEW_BATCH_BITRIX_SUBSET")
        preview = result.metadata.get("bitrix_batch_subset_preview") or {}
        self.assertEqual(preview.get("count"), 2)
        self.assertEqual(set(preview.get("selected_skus") or []), {SUBSET_NEW_SKU_1, SUBSET_NEW_SKU_2})
        self.assertIn(SUBSET_NEW_SKU_1, result.text)
        self.assertIn(SUBSET_NEW_SKU_2, result.text)
        self.assertNotIn("Строк было:", result.text)
        self.assertEqual(len(store.catalog(TENANT)), before_catalog_size)
        self.assertFalse(result.metadata.get("mutated"))

    async def test_fresh_attachment_two_named_skus_previews_both_with_managed_agent_enabled(self):
        await self._assert_direct_multi_preview(managed_enabled=True, request_suffix="managed-on")

    async def test_fresh_attachment_two_named_skus_previews_both_with_managed_agent_disabled(self):
        await self._assert_direct_multi_preview(managed_enabled=False, request_suffix="managed-off")

    async def test_one_product_postprocessing_failure_does_not_abort_other_selected_product(self):
        bridge, store = _bitrix_bridge_and_store()
        panda, artifact_service = _panda(bridge)
        ref = await _register_upload(artifact_service, content=_tcl_subset_two_new_no_title_bytes())
        before_catalog_size = len(store.catalog(TENANT))

        import business_assistant.product_enrichment_bridge as enrichment_bridge
        original_serialize = enrichment_bridge.serialize_write_request
        calls = {"n": 0}

        def _flaky_serialize(request):
            calls["n"] += 1
            if calls["n"] == 1:
                raise RuntimeError("synthetic_first_row_postprocess_failure")
            return original_serialize(request)

        with mock.patch.object(enrichment_bridge, "serialize_write_request", side_effect=_flaky_serialize):
            result = await panda.respond(
                ConversationRequest(
                    text=(
                        f"Подготовь для сайта только товары {SUBSET_NEW_SKU_1} и {SUBSET_NEW_SKU_2}. "
                        "Покажи полный предпросмотр того, что будет записано в Bitrix. "
                        "Пока ничего не записывай."
                    ),
                    tenant_id=TENANT,
                    user_id=OWNER,
                    request_id="direct-multi-row-isolation",
                    conversation_id=CONV,
                    attachment_refs=(ref,),
                )
            )

        self.assertEqual(result.metadata.get("action_decision"), "PREVIEW_BATCH_BITRIX_SUBSET")
        preview = result.metadata.get("bitrix_batch_subset_preview") or {}
        self.assertEqual(preview.get("count"), 1)
        self.assertEqual(len(preview.get("failed") or []), 1)
        self.assertIn(SUBSET_NEW_SKU_1, result.text)
        self.assertIn(SUBSET_NEW_SKU_2, result.text)
        self.assertIn("НЕ включён в список на создание", result.text)
        self.assertEqual(len(store.catalog(TENANT)), before_catalog_size)
        self.assertFalse(result.metadata.get("mutated"))
    async def test_fresh_attachment_clears_stale_frozen_batch_state(self):
        bridge, _store = _bitrix_bridge_and_store()
        panda, artifact_service = _panda(bridge)

        # First source establishes stale batch approval state.
        await _upload_and_analyze(panda, artifact_service)
        await panda.respond(
            ConversationRequest(
                text=BATCH_CHECK_TEXT, tenant_id=TENANT, user_id=OWNER,
                request_id="stale-r1", conversation_id=CONV,
            )
        )
        task = panda._action_store.get(tenant_id=TENANT, owner_id=OWNER, conversation_id=CONV)  # noqa: SLF001
        self.assertTrue(task.parameters.get("bitrix_batch_ready_rows"))

        # A new source must atomically reset all source-derived approval/
        # enrichment state before any selection is evaluated.
        ref = await _register_upload(artifact_service, content=_tcl_subset_two_new_no_title_bytes())
        await panda.respond(
            ConversationRequest(
                text="Вот новый прайс, проанализируй.",
                tenant_id=TENANT, user_id=OWNER, request_id='stale-r2',
                conversation_id=CONV, attachment_refs=(ref,),
            )
        )
        task = panda._action_store.get(tenant_id=TENANT, owner_id=OWNER, conversation_id=CONV)  # noqa: SLF001
        self.assertFalse(task.parameters.get("bitrix_batch_ready_rows"))
        self.assertFalse(task.parameters.get("bitrix_batch_all_ready_rows"))
        self.assertFalse(task.parameters.get("bitrix_batch_selected_skus"))
        self.assertFalse(task.parameters.get("bitrix_enrichment_write_request"))
        self.assertFalse(task.parameters.get("bitrix_product_fields"))


class BitrixBatchSubsetPreviewAcceptanceTests(unittest.IsolatedAsyncioTestCase):
    """34 NEW -> choose a named subset -> preview -> confirm only subset.
    The subset handoff must never mutate/filter the canonical Excel data.
    """

    async def _upload_two_new(self, panda, artifact_service):
        ref = await _register_upload(artifact_service, content=_tcl_subset_two_new_no_title_bytes())
        await panda.respond(
            ConversationRequest(
                text="Вот прайс-лист, проанализируй.",
                tenant_id=TENANT,
                user_id=OWNER,
                request_id="subset-r1",
                conversation_id=CONV,
                attachment_refs=(ref,),
            )
        )
        await panda.respond(
            ConversationRequest(
                text="Проанализируй весь прайс.",
                tenant_id=TENANT,
                user_id=OWNER,
                request_id="subset-r2",
                conversation_id=CONV,
            )
        )

    async def test_named_subset_is_previewed_without_excel_filter_and_only_subset_is_created(self):
        bridge, store = _bitrix_bridge_and_store()
        panda, artifact_service = _panda(bridge)
        await self._upload_two_new(panda, artifact_service)

        batch = await panda.respond(
            ConversationRequest(
                text=BATCH_CHECK_TEXT,
                tenant_id=TENANT,
                user_id=OWNER,
                request_id="subset-r3",
                conversation_id=CONV,
            )
        )
        check = batch.metadata.get("bitrix_existence_check") or {}
        self.assertEqual({r["sku"] for r in check.get("new") or []}, {SUBSET_NEW_SKU_1, SUBSET_NEW_SKU_2})
        before_catalog_size = len(store.catalog(TENANT))

        subset = await panda.respond(
            ConversationRequest(
                text=(
                    f"Подготовь для сайта только товары {SUBSET_NEW_SKU_1} и {SUBSET_NEW_SKU_2}. "
                    "Покажи полный предпросмотр того, что будет записано в Bitrix. "
                    "Пока ничего не записывай."
                ),
                tenant_id=TENANT,
                user_id=OWNER,
                request_id="subset-r4",
                conversation_id=CONV,
            )
        )

        self.assertEqual(subset.metadata.get("action_decision"), "PREVIEW_BATCH_BITRIX_SUBSET")
        self.assertFalse(subset.metadata.get("mutated"))
        self.assertNotIn("Строк было:", subset.text)
        preview = subset.metadata.get("bitrix_batch_subset_preview") or {}
        self.assertEqual(set(preview.get("selected_skus") or []), {SUBSET_NEW_SKU_1, SUBSET_NEW_SKU_2})
        self.assertEqual(preview.get("count"), 2)
        self.assertEqual(len(store.catalog(TENANT)), before_catalog_size)

        task = panda._action_store.get(tenant_id=TENANT, owner_id=OWNER, conversation_id=CONV)  # noqa: SLF001
        frozen = list(task.parameters.get("bitrix_batch_ready_rows") or [])
        self.assertEqual(len(frozen), 2)
        self.assertEqual({r.get("sku") for r in frozen}, {SUBSET_NEW_SKU_1, SUBSET_NEW_SKU_2})
        self.assertTrue(all(r.get("title") for r in frozen))
        self.assertTrue(all(r.get("prepared_write_request") for r in frozen))

        confirm = await panda.respond(
            ConversationRequest(
                text=BATCH_CONFIRM_TEXT,
                tenant_id=TENANT,
                user_id=OWNER,
                request_id="subset-r5",
                conversation_id=CONV,
            )
        )
        result = confirm.metadata.get("bitrix_batch_create_result") or {}
        self.assertEqual(result.get("total"), 2)
        self.assertEqual(result.get("created"), 2)
        self.assertEqual(len(store.catalog(TENANT)), before_catalog_size + 2)

        rerun = await panda.respond(
            ConversationRequest(
                text=BATCH_CHECK_TEXT,
                tenant_id=TENANT,
                user_id=OWNER,
                request_id="subset-r6",
                conversation_id=CONV,
            )
        )
        rerun_check = rerun.metadata.get("bitrix_existence_check") or {}
        self.assertEqual({r["sku"] for r in rerun_check.get("existing") or []}, {SUBSET_NEW_SKU_1, SUBSET_NEW_SKU_2})
        self.assertEqual(rerun_check.get("new") or [], [])


class TclStyleBatchWriteAndRerunAcceptanceTests(unittest.IsolatedAsyncioTestCase):
    """Mandatory acceptance #7, #9, #10, #11: the explicit batch
    confirmation creates ONLY the frozen READY_TO_CREATE rows, inactive,
    with a verified read-back, and a rerun of the SAME batch (either the
    stale confirmation itself or a fresh preview) never creates a
    duplicate."""

    async def _preview(self, panda):
        return await panda.respond(
            ConversationRequest(
                text=BATCH_CHECK_TEXT,
                tenant_id=TENANT,
                user_id=OWNER,
                request_id="r3",
                conversation_id=CONV,
            )
        )

    async def _confirm(self, panda, request_id: str):
        return await panda.respond(
            ConversationRequest(
                text=BATCH_CONFIRM_TEXT,
                tenant_id=TENANT,
                user_id=OWNER,
                request_id=request_id,
                conversation_id=CONV,
            )
        )

    async def test_confirmation_creates_only_new_row_inactive_with_verified_readback(self):
        bridge, store = _bitrix_bridge_and_store()
        panda, artifact_service = _panda(bridge)
        await _upload_and_analyze(panda, artifact_service)
        before_catalog_size = len(store.catalog(TENANT))

        preview = await self._preview(panda)
        self.assertEqual((preview.metadata.get("bitrix_existence_check") or {}).get("total"), 5)

        confirm = await self._confirm(panda, "r4")
        self.assertEqual(confirm.metadata.get("action_decision"), CONFIRM_BATCH_BITRIX_CREATE)
        self.assertTrue(confirm.metadata.get("mutated"))

        batch_result = confirm.metadata.get("bitrix_batch_create_result") or {}
        self.assertEqual(batch_result.get("created"), 1)
        self.assertEqual(batch_result.get("total"), 1)  # only the ONE frozen READY_TO_CREATE row
        self.assertIn(NEW_SKU, confirm.text)
        self.assertIn("создан", confirm.text.casefold())

        # #7/#9: exactly ONE product created (the NEW row) -- never the
        # EXISTING/AMBIGUOUS/INVALID rows.
        self.assertEqual(len(store.catalog(TENANT)), before_catalog_size + 1)

        # #10: read-back -- the created product is genuinely readable,
        # inactive, with the exact title from the price list.
        created_matches = bridge.check_live_existence(tenant_id=TENANT, sku=NEW_SKU)
        self.assertEqual(len(created_matches), 1)
        created_bitrix_id = str(created_matches[0].get("id") or created_matches[0].get("external_product_id"))
        self.assertTrue(created_bitrix_id)
        read_back = bridge.read_product(tenant_id=TENANT, bitrix_id=created_bitrix_id)
        self.assertEqual(read_back.get("name"), NEW_TITLE)
        self.assertFalse(read_back.get("active"))

    async def test_repeat_confirmation_of_same_frozen_batch_creates_no_duplicate(self):
        # #11 (rerun without duplicates), first shape: the user repeats
        # the EXACT SAME confirmation message without re-running the
        # preview first (the task's frozen ``bitrix_batch_ready_rows``
        # list is unchanged) -- the pre-create live re-check inside
        # ``_confirm_batch_bitrix_create`` must still refuse to create a
        # second product for the SAME SKU.
        bridge, store = _bitrix_bridge_and_store()
        panda, artifact_service = _panda(bridge)
        await _upload_and_analyze(panda, artifact_service)
        await self._preview(panda)

        first = await self._confirm(panda, "r4")
        self.assertEqual((first.metadata.get("bitrix_batch_create_result") or {}).get("created"), 1)
        after_first_create = len(store.catalog(TENANT))

        second = await self._confirm(panda, "r5")
        second_result = second.metadata.get("bitrix_batch_create_result") or {}
        self.assertEqual(second_result.get("created"), 0)
        self.assertEqual(second_result.get("skipped"), 1)
        self.assertIn("уже существ", second.text.casefold())
        self.assertEqual(len(store.catalog(TENANT)), after_first_create, "no second product must ever be created")

    async def test_repeat_whole_batch_operation_reclassifies_created_row_as_existing(self):
        # #11, second shape: the user reruns the WHOLE operation from the
        # top (a fresh preview after the earlier confirmation) -- the
        # already-created row must now show as EXISTING, never NEW again.
        bridge, store = _bitrix_bridge_and_store()
        panda, artifact_service = _panda(bridge)
        await _upload_and_analyze(panda, artifact_service)
        await self._preview(panda)
        await self._confirm(panda, "r4")
        after_first_create = len(store.catalog(TENANT))

        rerun_preview = await self._preview(panda)
        check = rerun_preview.metadata.get("bitrix_existence_check") or {}
        self.assertEqual({row["sku"] for row in check.get("new") or []}, set())
        self.assertIn(NEW_SKU, {row["sku"] for row in check.get("existing") or []})

        # With nothing left READY_TO_CREATE, the resolver fails closed
        # (no batch-create result at all -- distinct from, but equally
        # safe as, the "skipped" shape asserted in the previous test)
        # rather than ever attempting a second create for this SKU.
        rerun_confirm = await self._confirm(panda, "r5")
        rerun_result = rerun_confirm.metadata.get("bitrix_batch_create_result") or {}
        self.assertEqual(rerun_result.get("created") or 0, 0)
        self.assertEqual(len(store.catalog(TENANT)), after_first_create)


class SingleProductLiveDuplicateGuardAcceptanceTests(unittest.IsolatedAsyncioTestCase):
    """Mandatory acceptance #7 for the SINGLE-product conversational write
    path (Step 3's guard "must protect single-product create" too, reused
    from the SAME ``BitrixProductBridge.check_live_existence``): an
    explicit confirmation for a product whose SKU already exists live
    must be blocked, never create a duplicate."""

    async def test_exact_duplicate_single_product_create_is_blocked(self):
        bridge, store = _bitrix_bridge_and_store()
        panda, artifact_service = _panda(bridge)
        ref = await _register_upload(
            artifact_service,
            content=_tcl_xlsx_bytes(),
        )
        before_catalog_size = len(store.catalog(TENANT))

        await panda.respond(
            ConversationRequest(
                text=(
                    f"Найди товар {EXISTING_ACTIVE_SKU} и подготовь его для добавления в Bitrix. "
                    "Розничная цена 49990 \u20bd."
                ),
                tenant_id=TENANT,
                user_id=OWNER,
                request_id="r1",
                conversation_id=CONV,
                attachment_refs=(ref,),
            )
        )

        confirmation = await panda.respond(
            ConversationRequest(
                text=f"Подтверждаю: создай этот товар в Bitrix. Розничная цена 49990 \u20bd.",
                tenant_id=TENANT,
                user_id=OWNER,
                request_id="r2",
                conversation_id=CONV,
            )
        )

        self.assertEqual(confirmation.metadata.get("action_decision"), CALL_CONTROLLED_BITRIX_WRITE)
        self.assertEqual(
            confirmation.metadata.get("write_confirmation_event"), "WRITE_BLOCKED_LIVE_DUPLICATE_GUARD"
        )
        self.assertIn("уже существует", confirmation.text)
        self.assertEqual(len(store.catalog(TENANT)), before_catalog_size, "no duplicate must ever be created")


if __name__ == "__main__":
    unittest.main()
