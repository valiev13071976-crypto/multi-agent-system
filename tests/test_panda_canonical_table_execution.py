"""PANDA -- CANONICAL TABLE EXECUTION (NL -> ONE model semantic call ->
validated structured operation -> existing deterministic executor).

PRODUCTION DEFECT (after PR #91 fixed canonical Workset ownership): a
real PANDA_MANAGED_AGENT_ENABLED=true production request --

    "Увеличь розничную цену ... для всех товаров и покажи результат.
    Ничего в Bitrix не записывай."

-- did NOT execute the table transformation. Panda saw the data and even
computed aggregate stats, but answered conversationally ("Я не могу тут
массово пересчитать и вывести все строки...") instead of ever reaching
the EXISTING deterministic Data Intelligence executor.

ROOT CAUSE: the managed-agent integration boundary
(managed_agent_poc.runtime_subprocess) exposes exactly 3 READ-ONLY tools
-- analyze_spreadsheet / select_product / explain_bitrix_write_plan --
none of which can perform a bulk/structural table transform.

PR #92 (first fix): reused data_intel.nl_ops.compile_request (a bounded
deterministic RU/EN regex/stem compiler) as the semantic boundary. That
fix worked for a single, simple percent adjustment, but compile_request
cannot represent a COMPOUND request (scope A -> transform A, remainder ->
transform B, e.g. "первым трём +7%, остальным +15%") -- see
CompoundScopedOperationContractGapAuditTests below -- and its loose
_KEEP_ONLY_RE/price-filter grammar produced false positives on unrelated
free text, requiring four extra precedence-guard predicates that were
themselves more of the same regex-arbitration architecture.

PR #92 CORRECTION: compile_request is no longer the semantic boundary for
canonical table execution at all. Instead, ONE existing model-call seam
already available in this repo -- agents.openai_agent.OpenAIAgent.run (a
single one-shot HTTP call, already configured via the SAME
OPENAI_API_KEY/OPENAI_MODEL env vars managed_agent_poc requires) --
interprets the user's text EXACTLY ONCE per turn (see
data_intel.nl_plan_llm.compile_request_via_model) into STRICT JSON, which
is then deterministically parsed and validated into an OperationPlan
(with a minimal, generic scope extension -- all/row_range/remainder --
so ONE plan can now express a compound, multi-scope request). That plan
is executed by the SAME existing deterministic executor,
data_intel.transform.execute_plan.

PR #93 (production defect closure): PR #92 was deployed, but a REAL
production request over a column literally named
"Предоплата, Цена с НДС" (containing an internal comma) still failed --
Railway logs proved the OpenAI model call itself succeeded (HTTP 200),
yet the turn fell through to managed-agent product selection/enrichment
anyway. Reproduced here against a REAL, unmocked model call: the model
could not reliably reproduce that exact column string verbatim (it
echoed back only "Предоплата", truncated at the internal comma), so
column validation failed -- and PR #92's ``compile_request_via_model``
folded THAT validation failure into the exact same exception used for a
genuine "not a table operation" model judgment, which the caller could
not tell apart from a technical failure, so it silently deferred to the
managed agent. Two closures, both exercised below:

    1. the model is now asked for a stable ``column_id`` (``c0``, ``c1``,
       ...), never a raw column name, so it never has to reproduce an
       unusual string at all (see ``data_intel.nl_plan_llm._column_ids``);
    2. ``ModelPlanError`` now carries a typed ``status`` (``NOT_
       APPLICABLE`` vs ``MODEL_ERROR``/``PARSE_ERROR``/
       ``VALIDATION_ERROR``), and
       ``WorkflowPandaConversationGateway._maybe_execute_canonical_table_
       operation`` fails CLOSED (a plain "not executed" message) for any
       status other than ``OK``/``NOT_APPLICABLE`` instead of returning
       ``None`` -- so a technical failure of this boundary can never
       again be silently reinterpreted as an unrelated managed-agent
       workflow.

FIX: WorkflowPandaConversationGateway._maybe_execute_canonical_table_
operation (see business_assistant/conversation_gateway.py) is tried
BEFORE the managed-agent boundary for every eligible turn. It gives the
SAME data.excel_assistant/assist tool call the text FIRST, with a
use_model_plan=True flag so DataIntelToolAdapter routes to
DataIntelligenceService.execute_structured_plan_via_model instead of
execute_nl_request. A result is used (and the managed-agent boundary is
skipped entirely for this turn) when that tool's own status == "OK" --
i.e. the model judged this a genuine table operation AND its output
passed strict validation AND execute_plan ACTUALLY executed it. A
status == "NOT_APPLICABLE" (a genuine, validly-parsed non-table
judgment) returns None with zero side effects, so the managed agent
keeps handling every other conversational/product turn exactly as
before. Any OTHER status fails closed (PR #93) instead of falling
through.

No new agent/router/dataset store/executor is introduced: this reuses the
SAME one-shot model seam, the SAME (minimally extended) structured
contract, the SAME deterministic executor, and the SAME canonical
Workset (PR #91) every other FAMILY_EXCEL turn already uses.
"""

from __future__ import annotations

import io
import json
import os
import tempfile
import unittest
from unittest import mock

from business_assistant import workset as workset_lib
from managed_agent_poc.adapter import ManagedAgentPOC
from managed_agent_poc.panda_bridge import ENABLED_ENV_VAR
from tests.test_panda_canonical_workset_single_data_ownership import (
    EAN_A,
    EAN_B,
    EAN_C,
    EAN_D,
    FILENAME,
    NAME_A,
    NAME_B,
    NAME_C,
    NAME_D,
    OWNER,
    PURCHASE_A,
    PURCHASE_B,
    PURCHASE_C,
    PURCHASE_D,
    RETAIL_A,
    RETAIL_B,
    RETAIL_C,
    RETAIL_D,
    SKU_A,
    SKU_B,
    SKU_C,
    SKU_D,
    TENANT,
    _first_dataset_store,
    _xlsx_bytes,
)
from tests.test_panda_managed_agent_enrichment_delegation import (
    PRODUCTION_TEXT,
    _make_fake_run_turn,
    _raw_tool_fields,
)
from tests.test_panda_product_enrichment_conversational import _panda, _register_upload
from business_assistant.conversation_gateway import ConversationRequest

CONVERSATION_ID = "conv-table-exec"

RETAIL_COLUMN = "розница"
# Column order written by _xlsx_bytes (see
# tests/test_panda_canonical_workset_single_data_ownership.py):
# sku, product_name, category, brand, ean, purchase_price, розница.
RETAIL_COLUMN_ID = "c6"

# The exact production column name (PR #93 defect): contains an internal
# comma, which is exactly what a model can fail to reproduce verbatim.
PRODUCTION_COLUMN = "Предоплата, Цена с НДС"
PRODUCTION_COLUMN_ID = "c6"


def _analyze_plan_entry() -> dict:
    return {
        "tool_calls": [
            {
                "tool": "analyze_spreadsheet",
                "output": {"row_count": 4, "column_count": 7},
            }
        ],
        "final_output": "В таблице 4 строки.",
    }


def _select_product_plan_entry(*, identifier, name, purchase_price, retail_price) -> dict:
    return {
        "current_identifier": identifier,
        "tool_calls": [
            {
                "tool": "select_product",
                "output": {
                    "status": "SELECTED",
                    "matched_by": "next_unspecified",
                    "name": name,
                    "sku": identifier,
                    "category": "Телевизоры",
                    "brand": "LG",
                    "purchase_price": purchase_price,
                    "retail_price": retail_price,
                },
            }
        ],
        "final_output": "Выбран товар " + name + ".",
    }


def _tracking_fake_run_turn(plan):
    fake = _make_fake_run_turn(plan)
    calls = []

    def wrapper(self, **kwargs):
        calls.append(kwargs)
        return fake(self, **kwargs)

    return wrapper, calls


def _mock_model_json(payload):
    from agents.provider_result import ProviderResult

    async def _fake_run(self, prompt):
        return ProviderResult(text=json.dumps(payload), provider_id="openai", model_id="test-model")

    env_patch = mock.patch.dict(
        os.environ, {"OPENAI_API_KEY": "test-key", "OPENAI_MODEL": "test-model"}
    )
    model_patch = mock.patch("agents.openai_agent.OpenAIAgent.run", new=_fake_run)
    return env_patch, model_patch


def _real_credentials_available() -> bool:
    return bool((os.environ.get("OPENAI_API_KEY") or "").strip())


def _production_xlsx_bytes() -> bytes:
    """SAME 4-row fixture shape as _xlsx_bytes, but with the EXACT
    production column name from the PR #93 defect report (containing an
    internal comma) instead of the generic "розница" column."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["sku", "product_name", "category", "brand", "ean", "purchase_price", PRODUCTION_COLUMN])
    ws.append([SKU_A, NAME_A, "Телевизоры", "LG", EAN_A, PURCHASE_A, RETAIL_A])
    ws.append([SKU_B, NAME_B, "Телевизоры", "LG", EAN_B, PURCHASE_B, RETAIL_B])
    ws.append([SKU_C, NAME_C, "Телевизоры", "LG", EAN_C, PURCHASE_C, RETAIL_C])
    ws.append([SKU_D, NAME_D, "Телевизоры", "LG", EAN_D, PURCHASE_D, RETAIL_D])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class CanonicalTableExecutionAcceptanceTests(unittest.IsolatedAsyncioTestCase):

    async def asyncSetUp(self):
        self.tmp = tempfile.mkdtemp()
        self._old_data_dir = os.environ.get("PANDA_DATA_DIR")
        self._old_flag = os.environ.get(ENABLED_ENV_VAR)
        os.environ["PANDA_DATA_DIR"] = self.tmp
        os.environ[ENABLED_ENV_VAR] = "true"
        self.panda, self.artifact_service = _panda()
        self.artifact_id = await _register_upload(
            self.artifact_service,
            tenant=TENANT,
            owner=OWNER,
            conv=CONVERSATION_ID,
            filename=FILENAME,
            content=_xlsx_bytes(),
        )

    async def asyncTearDown(self):
        import shutil

        if self._old_data_dir is None:
            os.environ.pop("PANDA_DATA_DIR", None)
        else:
            os.environ["PANDA_DATA_DIR"] = self._old_data_dir
        if self._old_flag is None:
            os.environ.pop(ENABLED_ENV_VAR, None)
        else:
            os.environ[ENABLED_ENV_VAR] = self._old_flag
        shutil.rmtree(self.tmp, ignore_errors=True)

    def _workset(self):
        task = self.panda._action_store.get(
            tenant_id=TENANT, owner_id=OWNER, conversation_id=CONVERSATION_ID
        )
        return workset_lib.get_workset(task)

    async def _run_three_turn_journey(self, *, turn3_text, request_prefix, model_payload):
        plan = [
            _analyze_plan_entry(),
            _select_product_plan_entry(
                identifier=SKU_A, name=NAME_A, purchase_price=PURCHASE_A, retail_price=RETAIL_A
            ),
        ]
        fake_run_turn, run_turn_calls = _tracking_fake_run_turn(plan)
        env_patch, model_patch = _mock_model_json(model_payload)

        with mock.patch.object(ManagedAgentPOC, "run_turn", new=fake_run_turn):
            r1 = await self.panda.respond(
                ConversationRequest(
                    text="Проанализируй этот прайс.",
                    tenant_id=TENANT,
                    user_id=OWNER,
                    request_id=request_prefix + "-1",
                    conversation_id=CONVERSATION_ID,
                    attachment_refs=(self.artifact_id,),
                )
            )
            r2 = await self.panda.respond(
                ConversationRequest(
                    text="Возьми первый товар из этого прайса.",
                    tenant_id=TENANT,
                    user_id=OWNER,
                    request_id=request_prefix + "-2",
                    conversation_id=CONVERSATION_ID,
                )
            )
            self.assertEqual(len(run_turn_calls), 2)
            w_single = self._workset()
            self.assertEqual(w_single.scope, workset_lib.SCOPE_SINGLE)
            self.assertEqual(w_single.selected_identifiers, (SKU_A,))

            with env_patch, model_patch:
                r3 = await self.panda.respond(
                    ConversationRequest(
                        text=turn3_text,
                        tenant_id=TENANT,
                        user_id=OWNER,
                        request_id=request_prefix + "-3",
                        conversation_id=CONVERSATION_ID,
                    )
                )
            self.assertEqual(len(run_turn_calls), 2)

        return r1, r2, r3

    async def test_mandatory_acceptance_table_wide_increase_after_single_product_selection(self):
        model_payload = {
            "kind": "table_operation",
            "wants_workbook": False,
            "operations": [
                {
                    "scope": {"kind": "all"},
                    "column_id": RETAIL_COLUMN_ID,
                    "operation": "percent_round",
                    "value": "10",
                }
            ],
        }
        r1, r2, r3 = await self._run_three_turn_journey(
            turn3_text=(
                "Увеличь розничную цену на 10% для всех товаров и покажи результат. "
                "Ничего в Bitrix не записывай."
            ),
            request_prefix="rep",
            model_payload=model_payload,
        )

        self.assertNotIn("Приложите файл", r3.text)
        self.assertNotIn("не могу", r3.text.casefold())

        w_final = self._workset()
        self.assertIsNotNone(w_final)
        self.assertTrue(w_final.source_dataset_id)
        source_id = w_final.source_dataset_id

        self.assertEqual(w_final.scope, workset_lib.SCOPE_FULL_DATASET)
        self.assertEqual(w_final.selected_identifiers, ())
        self.assertNotEqual(w_final.current_dataset_id, source_id)

        self.assertIn("Строк было: 4, стало: 4", r3.text)
        preview = r3.metadata.get("table_operation_preview") or {}
        self.assertEqual(preview.get("row_count_before"), 4)
        self.assertEqual(preview.get("row_count_after"), 4)
        changed_rows = preview.get("changed_rows") or []
        self.assertEqual(len(changed_rows), 4)
        first = changed_rows[0]
        self.assertEqual(first["operation"], "percent_round")
        self.assertEqual(first["percent"], "10")
        self.assertTrue(first["source_value"])
        self.assertTrue(first["resulting_value"])
        self.assertNotEqual(first["source_value"], first["resulting_value"])
        self.assertTrue(any(str(v) == SKU_A for v in first["row"].values()))

        self.assertTrue(r3.metadata.get("canonical_table_execution"))
        self.assertEqual(r3.metadata.get("action_decision"), "CALL_TOOL")

        self.assertIsNone(self.panda._bitrix_bridge)
        self.assertNotIn("MaxTurnsExceeded", r3.text)

        store = _first_dataset_store(self.panda)
        original_rows = store.get_rows(source_id, tenant_id=TENANT)
        self.assertEqual(len(original_rows), 4)

    async def test_second_semantically_different_wording_and_value_same_mechanism(self):
        model_payload = {
            "kind": "table_operation",
            "operations": [
                {
                    "scope": {"kind": "all"},
                    "column_id": RETAIL_COLUMN_ID,
                    "operation": "percent_round",
                    "value": "-8",
                }
            ],
        }
        r1, r2, r3 = await self._run_three_turn_journey(
            turn3_text="Снизь розничную цену на 8% по всем позициям, покажи, что получилось.",
            request_prefix="rep2",
            model_payload=model_payload,
        )

        self.assertNotIn("Приложите файл", r3.text)
        w_final = self._workset()
        self.assertEqual(w_final.scope, workset_lib.SCOPE_FULL_DATASET)
        self.assertEqual(w_final.selected_identifiers, ())

        preview = r3.metadata.get("table_operation_preview") or {}
        changed_rows = preview.get("changed_rows") or []
        self.assertEqual(len(changed_rows), 4)
        first = changed_rows[0]
        self.assertEqual(first["percent"], "-8")
        self.assertTrue(any(str(v) == SKU_A for v in first["row"].values()))
        self.assertIsNone(self.panda._bitrix_bridge)

    async def test_compound_two_scope_operation_single_model_call(self):
        model_payload = {
            "kind": "table_operation",
            "operations": [
                {
                    "scope": {"kind": "row_range", "start": 0, "end": 3},
                    "column_id": RETAIL_COLUMN_ID,
                    "operation": "percent_round",
                    "value": "7",
                },
                {
                    "scope": {"kind": "remainder"},
                    "column_id": RETAIL_COLUMN_ID,
                    "operation": "percent_round",
                    "value": "15",
                },
            ],
        }
        r1, r2, r3 = await self._run_three_turn_journey(
            turn3_text=(
                "Первым трём товарам увеличь розничную цену на 7%, "
                "остальным увеличь розничную цену на 15%, покажи результат."
            ),
            request_prefix="compound",
            model_payload=model_payload,
        )

        w_final = self._workset()
        self.assertEqual(w_final.scope, workset_lib.SCOPE_FULL_DATASET)

        preview = r3.metadata.get("table_operation_preview") or {}
        changed_rows = preview.get("changed_rows") or []
        self.assertEqual(len(changed_rows), 4)
        by_row_index = {i: row for i, row in enumerate(changed_rows)}
        for i in range(3):
            self.assertEqual(by_row_index[i]["percent"], "7")
            self.assertEqual(by_row_index[i]["scope"]["kind"], "row_range")
        self.assertEqual(by_row_index[3]["percent"], "15")
        self.assertEqual(by_row_index[3]["scope"]["kind"], "remainder")
        for row in changed_rows:
            self.assertNotEqual(row["source_value"], row["resulting_value"])

        self.assertTrue(r3.metadata.get("canonical_table_execution"))
        self.assertIsNone(self.panda._bitrix_bridge)
        self.assertNotIn("MaxTurnsExceeded", r3.text)

    async def test_compound_differently_worded_with_different_values_zero_code_change(self):
        model_payload = {
            "kind": "table_operation",
            "operations": [
                {
                    "scope": {"kind": "row_range", "start": 0, "end": 2},
                    "column_id": RETAIL_COLUMN_ID,
                    "operation": "percent_round",
                    "value": "20",
                },
                {
                    "scope": {"kind": "remainder"},
                    "column_id": RETAIL_COLUMN_ID,
                    "operation": "percent_round",
                    "value": "-5",
                },
            ],
        }
        r1, r2, r3 = await self._run_three_turn_journey(
            turn3_text=(
                "Первым двум позициям поднимите цену на 20%, "
                "а на оставшиеся сделайте скидку 5%, покажи что вышло."
            ),
            request_prefix="compound2",
            model_payload=model_payload,
        )

        preview = r3.metadata.get("table_operation_preview") or {}
        changed_rows = preview.get("changed_rows") or []
        self.assertEqual(len(changed_rows), 4)
        self.assertEqual(changed_rows[0]["percent"], "20")
        self.assertEqual(changed_rows[1]["percent"], "20")
        self.assertEqual(changed_rows[2]["percent"], "-5")
        self.assertEqual(changed_rows[3]["percent"], "-5")
        self.assertIsNone(self.panda._bitrix_bridge)

    async def test_model_marks_non_table_request_not_applicable_falls_back_to_managed_agent(self):
        plan = [_analyze_plan_entry(), _analyze_plan_entry()]
        fake_run_turn, run_turn_calls = _tracking_fake_run_turn(plan)
        env_patch, model_patch = _mock_model_json({"kind": "not_applicable"})

        with mock.patch.object(ManagedAgentPOC, "run_turn", new=fake_run_turn):
            await self.panda.respond(
                ConversationRequest(
                    text="Проанализируй этот прайс.",
                    tenant_id=TENANT,
                    user_id=OWNER,
                    request_id="na-1",
                    conversation_id=CONVERSATION_ID,
                    attachment_refs=(self.artifact_id,),
                )
            )
            self.assertEqual(len(run_turn_calls), 1)
            with env_patch, model_patch:
                result = await self.panda.respond(
                    ConversationRequest(
                        text="А что там по остальным столбцам?",
                        tenant_id=TENANT,
                        user_id=OWNER,
                        request_id="na-2",
                        conversation_id=CONVERSATION_ID,
                    )
                )
            self.assertEqual(len(run_turn_calls), 2)
            self.assertIsNone(result.metadata.get("canonical_table_execution"))

    async def test_malformed_model_response_fails_closed_not_product_enrichment(self):
        """PR #93 requirement 5, targeted acceptance B: a TECHNICAL
        failure of the model-plan boundary (here: an unknown column_id,
        i.e. a VALIDATION_ERROR) must fail closed with a user-safe "not
        executed" message -- it must NEVER be silently reinterpreted as
        product selection/enrichment (the managed agent must not even be
        invoked for this turn), and the Workset/data must stay exactly as
        they were before the attempt."""

        model_payload = {
            "kind": "table_operation",
            "operations": [
                {"scope": {"kind": "all"}, "column_id": "c99", "operation": "percent_round", "value": "7"}
            ],
        }
        r1, r2, r3 = await self._run_three_turn_journey(
            turn3_text="Увеличь розничную цену на 7% для всех товаров и покажи результат.",
            request_prefix="failclosed",
            model_payload=model_payload,
        )

        self.assertIn("не удалось выполнить операцию над таблицей", r3.text.casefold())
        self.assertFalse(r3.metadata.get("canonical_table_execution"))
        self.assertTrue(r3.metadata.get("table_operation_failed_closed"))
        self.assertEqual(r3.metadata.get("table_operation_failure_status"), "VALIDATION_ERROR")
        self.assertEqual(r3.metadata.get("table_operation_failure_reason"), "unknown_column_id")

        # Zero side effects: the Workset is EXACTLY where turn 2 left it
        # (still SINGLE-product scope) -- nothing was promoted/derived
        # from a failed attempt.
        w_final = self._workset()
        self.assertEqual(w_final.scope, workset_lib.SCOPE_SINGLE)
        self.assertEqual(w_final.selected_identifiers, (SKU_A,))
        self.assertIsNone(self.panda._bitrix_bridge)

    async def test_single_product_and_plain_analysis_turns_still_use_managed_agent_unaffected(self):
        plan = [
            _analyze_plan_entry(),
            _select_product_plan_entry(
                identifier=SKU_B, name="Телевизор B", purchase_price="95000", retail_price="139990"
            ),
        ]
        fake_run_turn, run_turn_calls = _tracking_fake_run_turn(plan)
        # In production, OPENAI_API_KEY/OPENAI_MODEL are always set, so
        # the model is always actually invoked; the realistic outcome for
        # this text is a genuine "kind": "not_applicable" judgment, NOT a
        # missing-credentials technical failure (which would now fail
        # closed instead of deferring -- see PR #93). Mock the model
        # accordingly so this test reflects real production behavior.
        env_patch, model_patch = _mock_model_json({"kind": "not_applicable"})
        with mock.patch.object(ManagedAgentPOC, "run_turn", new=fake_run_turn):
            await self.panda.respond(
                ConversationRequest(
                    text="Проанализируй этот прайс.",
                    tenant_id=TENANT,
                    user_id=OWNER,
                    request_id="reg-1",
                    conversation_id=CONVERSATION_ID,
                    attachment_refs=(self.artifact_id,),
                )
            )
            self.assertEqual(len(run_turn_calls), 1)
            with env_patch, model_patch:
                await self.panda.respond(
                    ConversationRequest(
                        text="Покажи второй товар из прайса.",
                        tenant_id=TENANT,
                        user_id=OWNER,
                        request_id="reg-2",
                        conversation_id=CONVERSATION_ID,
                    )
                )
            self.assertEqual(len(run_turn_calls), 2)

    async def test_production_enrichment_text_with_coincidental_keep_only_wording_still_uses_managed_agent(self):
        plan = [
            {
                "current_identifier": SKU_A,
                "tool_calls": [
                    {
                        "tool": "select_product",
                        "output": {
                            "status": "SELECTED",
                            "matched_by": "next_unspecified",
                            **_raw_tool_fields(
                                name=NAME_A, sku=SKU_A, ean="8806096796849",
                                purchase_price=PURCHASE_A, retail_price=RETAIL_A,
                            ),
                        },
                    }
                ],
                "final_output": "Товар " + NAME_A + " подготовлен.",
            }
        ]
        fake_run_turn, run_turn_calls = _tracking_fake_run_turn(plan)
        # Same rationale as above: with real credentials in production,
        # the model is genuinely invoked and would judge this enrichment
        # text "not_applicable" -- mock that explicitly rather than
        # relying on a missing-credentials technical failure.
        env_patch, model_patch = _mock_model_json({"kind": "not_applicable"})
        with mock.patch.object(ManagedAgentPOC, "run_turn", new=fake_run_turn):
            with env_patch, model_patch:
                result = await self.panda.respond(
                    ConversationRequest(
                        text=PRODUCTION_TEXT,
                        tenant_id=TENANT,
                        user_id=OWNER,
                        request_id="prod-enrich-1",
                        conversation_id=CONVERSATION_ID,
                        attachment_refs=(self.artifact_id,),
                    )
                )
        self.assertEqual(len(run_turn_calls), 1)
        self.assertEqual(result.metadata.get("action_decision"), "MANAGED_AGENT")
        self.assertIsNone(result.metadata.get("canonical_table_execution"))
        store = _first_dataset_store(self.panda)
        w = self._workset()
        rows = store.get_rows(w.current_dataset_id, tenant_id=TENANT)
        self.assertEqual(len(rows), 4)


@unittest.skipUnless(
    _real_credentials_available(), "OPENAI_API_KEY not available in this environment -- see test docstring"
)
class RealUnmockedModelAcceptanceTests(unittest.IsolatedAsyncioTestCase):
    """PR #93 mandatory acceptance, part A: the EXACT production-shaped
    request against a fixture containing the EXACT production column
    "Предоплата, Цена с НДС", using a REAL (unmocked) OpenAI model call --
    agents.openai_agent.OpenAIAgent.run is NEVER patched anywhere in this
    class. Skipped automatically if OPENAI_API_KEY is unavailable in this
    environment; when it runs, it is definitive proof against the actual
    production defect (a model call that returns real, sometimes messy
    output), not just a mocked contract."""

    async def asyncSetUp(self):
        self.tmp = tempfile.mkdtemp()
        self._old_data_dir = os.environ.get("PANDA_DATA_DIR")
        self._old_flag = os.environ.get(ENABLED_ENV_VAR)
        self._old_api_key = os.environ.get("OPENAI_API_KEY")
        self._old_model = os.environ.get("OPENAI_MODEL")
        os.environ["PANDA_DATA_DIR"] = self.tmp
        os.environ[ENABLED_ENV_VAR] = "true"
        # This sandbox's injected OPENAI_API_KEY carries a trailing
        # newline that httpx rejects outright as an illegal header value
        # -- a purely local artifact of how the secret is injected here
        # (production's own key is unaffected; Railway logs already
        # proved a real HTTP 200 there). Strip it for this real call.
        if self._old_api_key:
            os.environ["OPENAI_API_KEY"] = self._old_api_key.strip()
        os.environ["OPENAI_MODEL"] = os.environ.get("OPENAI_MODEL", "gpt-4o-mini").strip() or "gpt-4o-mini"
        self.panda, self.artifact_service = _panda()
        self.conv_id = CONVERSATION_ID + "-real"
        self.artifact_id = await _register_upload(
            self.artifact_service,
            tenant=TENANT,
            owner=OWNER,
            conv=self.conv_id,
            filename=FILENAME,
            content=_production_xlsx_bytes(),
        )

    async def asyncTearDown(self):
        import shutil

        if self._old_data_dir is None:
            os.environ.pop("PANDA_DATA_DIR", None)
        else:
            os.environ["PANDA_DATA_DIR"] = self._old_data_dir
        if self._old_flag is None:
            os.environ.pop(ENABLED_ENV_VAR, None)
        else:
            os.environ[ENABLED_ENV_VAR] = self._old_flag
        if self._old_api_key is None:
            os.environ.pop("OPENAI_API_KEY", None)
        else:
            os.environ["OPENAI_API_KEY"] = self._old_api_key
        if self._old_model is None:
            os.environ.pop("OPENAI_MODEL", None)
        else:
            os.environ["OPENAI_MODEL"] = self._old_model
        shutil.rmtree(self.tmp, ignore_errors=True)

    async def test_real_model_call_executes_production_shaped_compound_request(self):
        plan = [_analyze_plan_entry()]
        fake_run_turn, run_turn_calls = _tracking_fake_run_turn(plan)

        with mock.patch.object(ManagedAgentPOC, "run_turn", new=fake_run_turn):
            await self.panda.respond(
                ConversationRequest(
                    text="Проанализируй этот прайс.",
                    tenant_id=TENANT,
                    user_id=OWNER,
                    request_id="real-1",
                    conversation_id=self.conv_id,
                    attachment_refs=(self.artifact_id,),
                )
            )
            self.assertEqual(len(run_turn_calls), 1)

            result = await self.panda.respond(
                ConversationRequest(
                    text=(
                        "Для первых трёх строк увеличь цену в столбце "
                        "«Предоплата, Цена с НДС» на 7%, для остальных строк — на 15%. "
                        "Покажи результат. Ничего в Bitrix не записывай."
                    ),
                    tenant_id=TENANT,
                    user_id=OWNER,
                    request_id="real-2",
                    conversation_id=self.conv_id,
                )
            )
            # THE production defect: the managed-agent loop must NOT be
            # entered for this turn at all.
            self.assertEqual(len(run_turn_calls), 1)

        self.assertTrue(
            result.metadata.get("canonical_table_execution"),
            f"expected canonical table execution, got: {result.text!r} / {result.metadata!r}",
        )
        preview = result.metadata.get("table_operation_preview") or {}
        changed_rows = preview.get("changed_rows") or []
        self.assertEqual(len(changed_rows), 4)
        self.assertEqual(changed_rows[0]["scope"]["kind"], "row_range")
        self.assertEqual(changed_rows[1]["scope"]["kind"], "row_range")
        self.assertEqual(changed_rows[2]["scope"]["kind"], "row_range")
        self.assertEqual(changed_rows[3]["scope"]["kind"], "remainder")
        for row in changed_rows:
            self.assertNotEqual(row["source_value"], row["resulting_value"])

        task = self.panda._action_store.get(tenant_id=TENANT, owner_id=OWNER, conversation_id=self.conv_id)
        w_final = workset_lib.get_workset(task)
        self.assertEqual(w_final.scope, workset_lib.SCOPE_FULL_DATASET)
        self.assertTrue(w_final.source_dataset_id)
        self.assertNotEqual(w_final.current_dataset_id, w_final.source_dataset_id)
        self.assertIsNone(self.panda._bitrix_bridge)


class CompoundScopedOperationContractGapAuditTests(unittest.TestCase):

    def test_compile_request_only_ever_compiles_one_percent_rule_per_request(self):
        from data_intel.nl_ops import OP_PERCENT_ROUND, compile_request
        from data_intel.service import DataIntelligenceService
        from data_intel.store import InMemoryDatasetStore

        svc = DataIntelligenceService(InMemoryDatasetStore())
        ingested = svc.ingest(_xlsx_bytes(), filename=FILENAME, tenant_id=TENANT)
        desc = svc.store.get_dataset(ingested["dataset_id"], tenant_id=TENANT)
        table = desc.tables[0]

        plan = compile_request(
            "Первым трём товарам увеличь розничную цену на 7%, "
            "остальным увеличь розничную цену на 15%, покажи результат.",
            table,
        )
        percent_ops = [op for op in plan.operations if op.op == OP_PERCENT_ROUND]
        self.assertEqual(len(percent_ops), 1)
        self.assertEqual(percent_ops[0].params["percent"], "7")

    def test_operation_plan_now_expresses_compound_scoped_rules_and_execute_plan_applies_them(self):
        from data_intel.nl_ops import (
            OP_PERCENT_ROUND,
            OperationPlan,
            OperationScope,
            PlannedOperation,
            SCOPE_REMAINDER,
            SCOPE_ROW_RANGE,
        )
        from data_intel.transform import execute_plan
        from data_intel.service import DataIntelligenceService
        from data_intel.store import InMemoryDatasetStore

        svc = DataIntelligenceService(InMemoryDatasetStore())
        ingested = svc.ingest(_xlsx_bytes(), filename=FILENAME, tenant_id=TENANT)
        desc = svc.store.get_dataset(ingested["dataset_id"], tenant_id=TENANT)
        table = desc.tables[0]
        rows = svc.store.get_rows(ingested["dataset_id"], tenant_id=TENANT)
        self.assertEqual(len(rows), 4)

        plan = OperationPlan(
            operations=(
                PlannedOperation(
                    OP_PERCENT_ROUND,
                    {"column": "розница", "percent": "7", "round_mode": None, "round_to": None},
                    scope=OperationScope(kind=SCOPE_ROW_RANGE, start=0, end=3),
                ),
                PlannedOperation(
                    OP_PERCENT_ROUND,
                    {"column": "розница", "percent": "15", "round_mode": None, "round_to": None},
                    scope=OperationScope(kind=SCOPE_REMAINDER),
                ),
            ),
        )
        result = execute_plan(rows, table.columns, plan)
        self.assertEqual(result.row_count_before, 4)
        self.assertEqual(result.row_count_after, 4)
        self.assertEqual(len(result.row_changes), 4)
        for change in result.row_changes[:3]:
            self.assertEqual(change["params"]["percent"], "7")
            self.assertEqual(change["scope"]["kind"], SCOPE_ROW_RANGE)
        self.assertEqual(result.row_changes[3]["params"]["percent"], "15")
        self.assertEqual(result.row_changes[3]["scope"]["kind"], SCOPE_REMAINDER)

        from decimal import Decimal

        self.assertEqual(
            Decimal(result.rows[0]["розница"]),
            (Decimal(RETAIL_A) * Decimal("1.07")).quantize(Decimal("0.01")),
        )


if __name__ == "__main__":
    unittest.main()
